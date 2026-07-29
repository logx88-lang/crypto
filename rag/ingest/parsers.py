"""포맷별 파서 — 등장 순서 보존 추출, 표 마크다운 직렬화, doc_type 분류.

지원: xlsx / docx / pptx / pdf / txt  (hex/이미지는 별도 경로)
공통 원칙: 표준 양식 비의존. 스타일 이름에 의존하지 않고 텍스트·표를 범용 추출한다.
"""
from __future__ import annotations

import os
from typing import Optional

from ..config import CONFIG
from .models import Element, ParsedDoc
from .tables import table_to_markdown

SUPPORTED_EXTS = {"xlsx", "docx", "doc", "pptx", "pdf", "txt",
                  "png", "jpg", "jpeg", "bmp", "tiff", "tif"}   # 이미지=OCR(§8), doc=변환


# ---------------------------------------------------------------------------
def _classify(doc_title: str, sample_text: str) -> str:
    hay = (doc_title + "\n" + sample_text).lower()
    for kw in CONFIG.protocol_keywords:
        if kw.lower() in hay:
            return "protocol_spec"
    return "general"


def _title_from_name(path: str) -> str:
    return os.path.splitext(os.path.basename(path))[0]


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------
def parse_xlsx(path: str) -> ParsedDoc:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    elements: list = []
    sample = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        # 병합셀 fill-forward
        maxc = max((len(r) for r in rows), default=0)
        rows = [r + [None] * (maxc - len(r)) for r in rows]
        for rng in list(ws.merged_cells.ranges):
            tl = rows[rng.min_row - 1][rng.min_col - 1]
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if 0 <= r - 1 < len(rows) and 0 <= c - 1 < len(rows[r - 1]):
                        rows[r - 1][c - 1] = tl
        md = table_to_markdown(rows, header=True)
        if md:
            elements.append(Element("table", md, {"sheet_name": ws.title}))
            sample.append(md[:200])
    title = _title_from_name(path)
    return ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                     doc_type=_classify(title, "\n".join(sample)), ext="xlsx",
                     elements=elements)


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
def parse_docx(path: str) -> ParsedDoc:
    from docx import Document
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl
    from docx.text.paragraph import Paragraph
    from docx.table import Table

    doc = Document(path)
    elements: list = []
    section = ""
    first_heading = ""
    heading_stack: list = []   # [(level, text)] — 제목 계층 추적

    def _loc():
        loc = {"section": section}
        path = " > ".join(t for _l, t in heading_stack)
        if path:
            loc["section_path"] = path   # 상위→하위 제목 경로(표 검색 노출용)
        return loc

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            p = Paragraph(child, doc)
            txt = (p.text or "").strip()
            if not txt:
                continue
            style = (p.style.name or "") if p.style else ""
            if style.startswith("Heading") or style.startswith("Title"):
                tail = style.split()[-1] if style.split() else ""
                lvl = int(tail) if tail.isdigit() else 1
                heading_stack = [(l, t) for (l, t) in heading_stack if l < lvl]
                heading_stack.append((lvl, txt))
                section = txt
                if not first_heading:
                    first_heading = txt
            elements.append(Element("text", txt, _loc()))
        elif isinstance(child, CT_Tbl):
            t = Table(child, doc)
            rows = [[c.text for c in row.cells] for row in t.rows]
            md = table_to_markdown(rows, header=True)
            if md:
                elements.append(Element("table", md, _loc()))
    title = first_heading or _title_from_name(path)
    sample = "\n".join(e.text[:200] for e in elements[:8])
    return ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                     doc_type=_classify(title, sample), ext="docx", elements=elements)


# ---------------------------------------------------------------------------
# doc (구형 이진 워드) — .docx 로 변환 후 파싱(표/제목계층 보존)
# ---------------------------------------------------------------------------
def _convert_doc_to_docx(path: str):
    """구형 .doc → 임시 .docx 변환. MS Word(win32com) 우선, 없으면 LibreOffice.

    반환: (docx경로 또는 None, 실패사유 리스트). 사유는 왜 안 됐는지 UI에 노출용.
    """
    import tempfile
    abspath = os.path.abspath(path)
    out = os.path.abspath(os.path.join(tempfile.mkdtemp(), "converted.docx"))
    reasons = []
    # 1) MS Word COM (Windows + Word 설치 시)
    try:
        import win32com.client  # type: ignore
        # 웹서버 스레드풀/Streamlit 스레드에서 COM 사용 시 스레드별 초기화 필수 —
        # 없으면 'CoInitialize has not been called' 로 Word가 있어도 실패한다.
        _com = None
        try:
            import pythoncom  # type: ignore
            pythoncom.CoInitialize()
            _com = pythoncom
        except Exception:
            pass
        try:
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = False
            try:
                d = word.Documents.Open(abspath, ReadOnly=True, AddToRecentFiles=False)
                d.SaveAs2(out, FileFormat=16)   # 16 = wdFormatDocumentDefault(.docx)
                d.Close(False)
            finally:
                word.Quit()
            if os.path.exists(out):
                return out, reasons
            reasons.append("Word 변환 결과 파일 미생성")
        except Exception as e:
            reasons.append(f"Word COM 실패({type(e).__name__}): {e}")
        finally:
            if _com is not None:
                try:
                    _com.CoUninitialize()
                except Exception:
                    pass
    except ImportError as e:
        # 원인 그대로 노출(진짜 미설치 vs DLL 로드 실패 구분). DLL 실패면
        # .venv\Scripts\pywin32_postinstall.py -install 로 해결되는 경우가 많다.
        reasons.append(f"pywin32(win32com) 임포트 실패: {e}")
    # 2) LibreOffice (soffice --headless)
    import subprocess
    outdir = os.path.dirname(out)
    lo_tried = False
    for exe in ("soffice", "soffice.exe",
                r"C:\Program Files\LibreOffice\program\soffice.exe"):
        try:
            subprocess.run([exe, "--headless", "--convert-to", "docx",
                            "--outdir", outdir, abspath],
                           check=True, timeout=180,
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            lo_tried = True
            cand = os.path.join(outdir,
                                os.path.splitext(os.path.basename(path))[0] + ".docx")
            if os.path.exists(cand):
                return cand, reasons
        except FileNotFoundError:
            continue
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
            reasons.append(f"LibreOffice 실패: {e}")
            lo_tried = True
    if not lo_tried:
        reasons.append("LibreOffice(soffice) 미설치")
    return None, reasons


def parse_doc(path: str) -> ParsedDoc:
    tmp, reasons = _convert_doc_to_docx(path)
    if not tmp:
        raise ValueError(
            "'.doc' 변환 실패 [" + " / ".join(reasons) + "]. "
            "MS Word 또는 LibreOffice가 필요합니다(없으면 .docx 로 저장 후 업로드). "
            "Word가 설치돼 있는데도 실패하면 pywin32 설치 여부를 확인하세요.")
    doc = parse_docx(tmp)      # 변환된 docx 를 기존 파서로(표·제목계층 그대로)
    # 원본 정보로 교체
    doc.source_file = os.path.basename(path)
    doc.ext = "doc"
    try:
        import shutil
        shutil.rmtree(os.path.dirname(tmp), ignore_errors=True)
    except Exception:
        pass
    return doc


# ---------------------------------------------------------------------------
# pptx
# ---------------------------------------------------------------------------
def parse_pptx(path: str) -> ParsedDoc:
    from pptx import Presentation
    prs = Presentation(path)
    elements: list = []
    sample = []
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                rows = [[tbl.cell(r, c).text for c in range(len(tbl.columns))]
                        for r in range(len(tbl.rows))]
                md = table_to_markdown(rows, header=True)
                if md:
                    elements.append(Element("table", md, {"slide_no": idx}))
            elif shape.has_text_frame:
                txt = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if txt:
                    elements.append(Element("text", txt, {"slide_no": idx}))
                    if len(sample) < 8:
                        sample.append(txt[:120])
    title = _title_from_name(path)
    return ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                     doc_type=_classify(title, "\n".join(sample)), ext="pptx",
                     elements=elements)


# ---------------------------------------------------------------------------
# pdf
# ---------------------------------------------------------------------------
def _norm_row(row) -> tuple:
    return tuple(str(c or "").strip() for c in row)


def _merge_cross_page_tables(page_tbls: dict, page_h: dict,
                             bottom_margin: float = 120.0,
                             top_margin: float = 150.0) -> dict:
    """페이지 걸침 표 병합 → {시작페이지: [(rows, page_span|None), ...]}.

    판단: 이전 페이지의 '마지막 표'가 페이지 하단(bottom_margin 이내)까지 닿고,
    다음 페이지의 '첫 표'가 상단(top_margin 이내)에서 시작하며 열 수가 같으면
    같은 표의 연속으로 보고 행을 이어붙인다(다음 조각의 반복 헤더 제거). 여러 페이지 연쇄 지원.
    """
    out: dict = {p: [] for p in page_tbls}
    open_chain = None      # {"start": p, "rows": [...], "ncols": n, "header": tuple, "end": p}
    for p in sorted(page_tbls):
        tbls = page_tbls[p]
        h = page_h.get(p, 842.0)
        for ti, t in enumerate(tbls):
            cont = (open_chain is not None and ti == 0
                    and t["top"] <= top_margin
                    and t["ncols"] == open_chain["ncols"]
                    and p == open_chain["end"] + 1)
            if cont:
                rows = t["rows"]
                # 다음 페이지 첫 행이 반복 헤더(첫 조각 헤더와 동일)면 제거
                if rows and _norm_row(rows[0]) == open_chain["header"]:
                    rows = rows[1:]
                open_chain["rows"].extend(rows)
                open_chain["end"] = p
            else:
                if open_chain is not None:
                    _close_chain(out, open_chain)
                open_chain = {"start": p, "end": p, "rows": list(t["rows"]),
                              "ncols": t["ncols"], "header": _norm_row(t["rows"][0])}
            # 페이지 하단까지 닿지 않거나 마지막 표가 아니면 다음 페이지로 이어질 수 없음
            is_last = ti == len(tbls) - 1
            if open_chain is not None and not (is_last and t["bottom"] >= h - bottom_margin):
                _close_chain(out, open_chain)
                open_chain = None
    if open_chain is not None:
        _close_chain(out, open_chain)
    return out


def _close_chain(out: dict, chain: dict) -> None:
    span = f"{chain['start']}-{chain['end']}" if chain["end"] > chain["start"] else None
    out.setdefault(chain["start"], []).append((chain["rows"], span))


def parse_pdf(path: str, ocr_backend=None) -> ParsedDoc:
    import pdfplumber
    from .ocr import ocr_image, get_ocr_backend
    from ..config import CONFIG

    backend = ocr_backend or get_ocr_backend()
    # 스캔 페이지를 실제로 만났을 때만 백엔드 가용성 확인(정상 PDF는 paddle import 회피).
    _ocr_state = {}
    def _ocr_ready():
        if "on" not in _ocr_state:
            _ocr_state["on"] = CONFIG.ocr_enabled and backend.available()
        return _ocr_state["on"]

    elements: list = []
    sample = []
    scanned_pages = []      # OCR 못한(비활성) 스캔 페이지
    ocr_pages = []          # OCR로 텍스트 복원한 페이지
    page_text: dict = {}    # idx -> 텍스트 Element
    page_tbls: dict = {}    # idx -> [{rows, top, bottom, ncols}] (페이지 걸침 병합 판단용)
    page_h: dict = {}
    with pdfplumber.open(path) as pdf:
        for idx, page in enumerate(pdf.pages, start=1):
            text = (page.extract_text() or "").strip()
            try:
                found = page.find_tables() or []
            except Exception as e:
                print(f"[경고] PDF {os.path.basename(path)} p{idx} 표 추출 실패(표 누락 가능): {e}",
                      flush=True)
                found = []
            tables = []
            for t in found:
                rows = t.extract() or []
                rows = [r for r in rows if any(c not in (None, "") for c in r)]
                if rows:
                    tables.append({"rows": rows, "top": float(t.bbox[1]),
                                   "bottom": float(t.bbox[3]),
                                   "ncols": max(len(r) for r in rows)})
            page_h[idx] = float(page.height)
            page_tbls[idx] = tables
            if not text and not tables:
                # 스캔(이미지) 페이지 → OCR 라우팅(§3 pdf)
                if _ocr_ready():
                    try:
                        pil = page.to_image(resolution=CONFIG.ocr_dpi).original
                        otext = (ocr_image(pil, backend=backend) or "").strip()
                    except Exception:
                        otext = ""
                    if otext:
                        page_text[idx] = Element("text", otext,
                                                 {"page_no": idx, "source": "ocr"})
                        ocr_pages.append(idx)
                        if len(sample) < 5:
                            sample.append(otext[:200])
                        continue
                scanned_pages.append(idx)
                continue
            if text:
                page_text[idx] = Element("text", text, {"page_no": idx})
                if len(sample) < 5:
                    sample.append(text[:200])

    # 페이지에 걸쳐 이어지는 표 병합(예: 명령 목록 표가 4~5페이지에 나뉜 경우 → 하나의 표).
    merged = _merge_cross_page_tables(page_tbls, page_h)
    # 원문 읽기 순서 유지: 페이지순으로 텍스트 → 그 페이지에서 시작한 표
    for idx in sorted(set(list(page_text.keys()) + list(merged.keys()))):
        if idx in page_text:
            elements.append(page_text[idx])
        for rows, span in merged.get(idx, []):
            md = table_to_markdown(rows, header=True)
            if md:
                loc = {"page_no": idx}
                if span:
                    loc["page_span"] = span      # 예: "4-5" (여러 페이지에 걸친 표)
                elements.append(Element("table", md, loc))
    title = _title_from_name(path)
    doc = ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                    doc_type=_classify(title, "\n".join(sample)), ext="pdf",
                    elements=elements)
    if scanned_pages:
        # OCR 비활성/미설치로 텍스트 못 뽑은 스캔 페이지 메모
        doc.elements.append(Element("text",
                                    f"[OCR 필요 페이지(백엔드 미설치): {scanned_pages}]",
                                    {"page_no": scanned_pages[0], "section": "_ocr_todo"}))
    return doc


# ---------------------------------------------------------------------------
# txt (인코딩 자동 판별)
# ---------------------------------------------------------------------------
def parse_txt(path: str) -> ParsedDoc:
    from charset_normalizer import from_path
    best = from_path(path).best()
    text = str(best) if best else ""
    enc = best.encoding if best else "unknown"
    title = _title_from_name(path)
    elements = [Element("text", text.strip(), {"section": "", "encoding": enc})] if text.strip() else []
    return ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                     doc_type=_classify(title, text[:400]), ext="txt", elements=elements)


# ---------------------------------------------------------------------------
# image (OCR — §8)
# ---------------------------------------------------------------------------
def parse_image(path: str, ocr_backend=None) -> ParsedDoc:
    from .ocr import ocr_image
    text = (ocr_image(path, backend=ocr_backend) or "").strip()
    title = _title_from_name(path)
    if text:
        elements = [Element("text", text, {"section": "", "source": "ocr"})]
        dtype = _classify(title, text[:400])
    else:
        # OCR 비활성/미설치/무텍스트 → 마커(인덱싱은 되되 검색가치 낮음 명시)
        elements = [Element("text", "[OCR 텍스트 없음 또는 OCR 백엔드 미설치]",
                            {"section": "_ocr_unavailable", "source": "ocr"})]
        dtype = "general"
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    return ParsedDoc(source_file=os.path.basename(path), doc_title=title,
                     doc_type=dtype, ext=ext, elements=elements)


_DISPATCH = {
    "xlsx": parse_xlsx, "docx": parse_docx, "doc": parse_doc, "pptx": parse_pptx,
    "pdf": parse_pdf, "txt": parse_txt,
    "png": parse_image, "jpg": parse_image, "jpeg": parse_image,
    "bmp": parse_image, "tiff": parse_image, "tif": parse_image,
}


def parse_file(path: str) -> ParsedDoc:
    ext = os.path.splitext(path)[1].lower().lstrip(".")
    if ext not in _DISPATCH:
        raise ValueError(f"지원하지 않는 확장자: .{ext} (지원: {sorted(SUPPORTED_EXTS)})")
    return _DISPATCH[ext](path)
