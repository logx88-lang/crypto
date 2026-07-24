"""근거 문서 원본 미리보기 → HTML 조각(웹 UI 모달용). Streamlit 비의존.

xlsx→HTML 표, docx/pptx→문단+표, pdf→해당 페이지 이미지(base64), txt→텍스트,
doc→docx 변환(Word/LibreOffice, 결과 캐시) 후 렌더.
excerpt(답변이 참고한 발췌)가 주어지면 그 부분을 노랗게 강조하고 자동 스크롤한다.
이미 설치된 라이브러리(openpyxl·pdfplumber·python-docx·python-pptx) 사용.
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import os
import re
import tempfile

from ..config import CONFIG

_MAX_ELEMENTS = 400
_HL_STYLE = "background:#fff3bf;border-radius:3px"     # 강조(연노랑)
_SCROLL_JS = ("<script>setTimeout(function(){var el=document.querySelector"
              "('#modal-body .pv-hl');if(el)el.scrollIntoView({block:'center'});},60)</script>")


def _esc(s):
    return html.escape(str(s if s is not None else ""))


def _norm(s) -> str:
    """공백 제거 정규화 — 파싱/직렬화 차이에 흔들리지 않는 부분 일치용."""
    return re.sub(r"\s+", "", str(s or ""))


def _hit(unit_text, excerpt_norm: str, min_len: int = 6) -> bool:
    """문단/행/줄 단위 텍스트가 발췌 안에 들어 있으면 강조 대상."""
    if not excerpt_norm:
        return False
    n = _norm(unit_text)
    return len(n) >= min_len and n in excerpt_norm


def _rows_to_html(rows, hl=None) -> str:
    hl = hl or set()
    if not rows:
        return "<p class='text-gray-400'>(빈 표)</p>"
    ncol = max((len(r) for r in rows), default=0)
    out = ["<table class='min-w-full text-sm border-collapse'>"]
    for ri, r in enumerate(rows):
        cells = [("" if v is None else str(v)) for v in r] + [""] * (ncol - len(r))
        tag = "th" if ri == 0 else "td"
        cls = ("bg-gray-100 font-semibold" if ri == 0 else "") + " border border-gray-200 px-2 py-1 text-left"
        mark = f" class='pv-hl' style='{_HL_STYLE}'" if ri in hl else ""
        out.append(f"<tr{mark}>" + "".join(f"<{tag} class='{cls}'>{_esc(c)}</{tag}>" for c in cells) + "</tr>")
    out.append("</table>")
    return "".join(out)


def render_file(rel_path: str, loc: dict = None, excerpt: str = None) -> str:
    loc = loc or {}
    if not rel_path:
        return "<p class='text-amber-600'>원본 경로 정보가 없습니다.</p>"
    # 경로 이탈(../) 차단 — data_dir 밖 파일 열람 방지(절대경로 기준 컨테인먼트)
    base = os.path.abspath(CONFIG.data_dir)
    path = os.path.abspath(os.path.join(base, *rel_path.replace("\\", "/").split("/")))
    if not path.startswith(base + os.sep) or not os.path.exists(path):
        return f"<p class='text-amber-600'>파일을 찾을 수 없습니다: {_esc(rel_path)}</p>"
    ext = rel_path.rsplit(".", 1)[-1].lower() if "." in rel_path else ""
    exn = _norm(excerpt) if excerpt else ""
    head = f"<div class='text-xs text-gray-500 mb-3'>📄 {_esc(rel_path)}" + \
           ("<span class='ml-2' style='" + _HL_STYLE + "'>&nbsp;노란 부분&nbsp;</span>= 답변이 참고한 발췌"
            if exn else "") + "</div>"
    try:
        if ext == "xlsx":
            body = _xlsx(path, loc, exn)
        elif ext == "pdf":
            body = _pdf(path, loc, exn)
        elif ext == "docx":
            body = _docx(path, exn)
        elif ext == "doc":
            body = _doc(path, exn)
        elif ext == "pptx":
            body = _pptx(path, loc, exn)
        elif ext in ("txt", "log", "csv"):
            body = _txt(path, exn)
        else:
            body = f"<p class='text-gray-500'>'{_esc(ext)}' 형식 미리보기 미지원.</p>"
        scroll = _SCROLL_JS if (exn and "pv-hl" in body) else ""
        return head + body + scroll
    except Exception as e:
        return head + f"<p class='text-amber-600'>미리보기 실패: {_esc(type(e).__name__)}: {_esc(e)}</p>"


def _xlsx(path, loc, exn=""):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = wb.sheetnames
    target = loc.get("sheet_name") if loc.get("sheet_name") in sheets else sheets[0]
    tabs = "".join(
        f"<button class='px-2 py-1 text-xs rounded mr-1 mb-1 {'bg-[#d6006a] text-white' if s==target else 'bg-gray-100'}' "
        f"hx-get='/preview?rel_path={_url(path_relparam(path))}&sheet_name={_url(s)}' "
        f"hx-target='#modal-body'>{_esc(s)}</button>" for s in sheets)
    rows = [[c.value for c in row] for row in wb[target].iter_rows()]
    # 행 전체 텍스트 또는 행 내 '긴 셀'이 발췌에 있으면 그 행 강조 (표 직렬화 방식과 무관하게 매칭)
    hl = set()
    for i, r in enumerate(rows):
        joined = "".join(str(v) for v in r if v is not None)
        if _hit(joined, exn, 8) or any(_hit(v, exn, 8) for v in r if v is not None):
            hl.add(i)
    return f"<div class='mb-2'>{tabs}</div><div class='overflow-auto max-h-[65vh]'>{_rows_to_html(rows, hl)}</div>"


def path_relparam(path):
    # data_dir 기준 상대경로 복원(시트 전환 링크용)
    rel = os.path.relpath(path, CONFIG.data_dir).replace("\\", "/")
    return rel


def _url(s):
    import urllib.parse
    return urllib.parse.quote(str(s))


def _pdf(path, loc, exn=""):
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        pno = min(max(int(loc.get("page_no") or 1), 1), n)
        page = pdf.pages[pno - 1]
        img = page.to_image(resolution=130)
        # 발췌에 나오는 단어 위치에 노랑 박스(가능할 때만; 실패해도 페이지는 표시)
        if exn:
            try:
                words = page.extract_words()
                rects = [w for w in words
                         if len(_norm(w.get("text"))) >= 3 and _norm(w.get("text")) in exn]
                if rects and len(rects) <= 200:
                    img.draw_rects(
                        [(w["x0"], w["top"], w["x1"], w["bottom"]) for w in rects],
                        stroke="#f1c40f", stroke_width=2, fill=(241, 196, 15, 60))
            except Exception:
                pass
        buf = io.BytesIO(); img.save(buf, format="PNG")
        b64 = base64.b64encode(buf.getvalue()).decode()
    rel = path_relparam(path)
    nav = ""
    if n > 1:
        prev = f"<button class='px-2 py-1 bg-gray-100 rounded text-xs' hx-get='/preview?rel_path={_url(rel)}&page_no={pno-1}' hx-target='#modal-body' {'disabled' if pno<=1 else ''}>◀ 이전</button>"
        nxt = f"<button class='px-2 py-1 bg-gray-100 rounded text-xs' hx-get='/preview?rel_path={_url(rel)}&page_no={pno+1}' hx-target='#modal-body' {'disabled' if pno>=n else ''}>다음 ▶</button>"
        nav = f"<div class='flex items-center gap-2 mb-2 text-xs text-gray-500'>{prev}<span>{pno} / {n}</span>{nxt}</div>"
    return f"{nav}<div class='overflow-auto max-h-[70vh]'><img class='w-full border border-gray-200' src='data:image/png;base64,{b64}'></div>"


def _docx(path, exn=""):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out, cnt = ["<div class='overflow-auto max-h-[70vh] space-y-2'>"], 0
    for ch in d.element.body.iterchildren():
        if cnt >= _MAX_ELEMENTS:
            out.append("<p class='text-gray-400 text-xs'>… (문서가 길어 이후 생략)</p>"); break
        if ch.tag.endswith("}p"):
            t = Paragraph(ch, d).text
            if t.strip():
                mark = f" class='pv-hl' style='{_HL_STYLE}'" if _hit(t, exn) else ""
                out.append(f"<p{mark}>{_esc(t)}</p>"); cnt += 1
        elif ch.tag.endswith("}tbl"):
            rows = [[c.text for c in r.cells] for r in Table(ch, d).rows]
            hl = {i for i, r in enumerate(rows) if _hit("".join(r), exn, 8)}
            out.append(_rows_to_html(rows, hl)); cnt += 1
    out.append("</div>")
    return "".join(out)


def _doc(path, exn=""):
    """구형 .doc — Word/LibreOffice로 docx 변환(파일별 캐시) 후 docx 렌더."""
    key = hashlib.sha1(f"{os.path.abspath(path)}|{os.path.getmtime(path)}".encode()).hexdigest()[:16]
    cache = os.path.join(tempfile.gettempdir(), "rag_preview_cache")
    os.makedirs(cache, exist_ok=True)
    cached = os.path.join(cache, key + ".docx")
    if not os.path.exists(cached):
        from ..ingest.parsers import _convert_doc_to_docx
        conv, reasons = _convert_doc_to_docx(path)
        if not conv:
            why = "; ".join(reasons) or "변환 실패"
            return ("<p class='text-amber-600'>.doc 변환 실패 — 서버에 MS Word 또는 LibreOffice가 "
                    f"필요합니다.<br><span class='text-xs'>{_esc(why)}</span></p>")
        import shutil
        shutil.copyfile(conv, cached)
    return _docx(cached, exn)


def _pptx(path, loc, exn=""):
    from pptx import Presentation
    prs = Presentation(path)
    target = None
    try:
        target = int(loc.get("slide_no") or 0)
    except Exception:
        pass
    out = ["<div class='overflow-auto max-h-[70vh] space-y-3'>"]
    for i, slide in enumerate(prs.slides, start=1):
        # 근거 슬라이드는 제목줄 자체를 강조(+스크롤 목표)
        tmark = f" class='pv-hl' style='{_HL_STYLE}'" if (target and i == target) else ""
        out.append(f"<div{tmark or ' class=\"font-semibold text-gray-600\"'}>— 슬라이드 {i} —</div>")
        for shape in slide.shapes:
            if shape.has_table:
                tbl = shape.table
                rows = [[tbl.cell(r, c).text for c in range(len(tbl.columns))]
                        for r in range(len(tbl.rows))]
                hl = {ri for ri, r in enumerate(rows) if _hit("".join(r), exn, 8)}
                out.append(_rows_to_html(rows, hl))
            elif shape.has_text_frame and shape.text_frame.text.strip():
                t = shape.text_frame.text
                mark = f" class='pv-hl' style='{_HL_STYLE}'" if _hit(t, exn) else ""
                out.append(f"<p{mark}>{_esc(t)}</p>")
    out.append("</div>")
    return "".join(out)


def _txt(path, exn=""):
    from charset_normalizer import from_path
    try:
        text = str(from_path(path).best())
    except Exception:
        text = open(path, encoding="utf-8", errors="replace").read()
    text = text[:20000]
    if not exn:
        return f"<pre class='whitespace-pre-wrap text-xs overflow-auto max-h-[70vh] bg-gray-50 p-3 rounded'>{_esc(text)}</pre>"
    # 줄 단위 강조 — 발췌에 포함된 줄에 <mark>
    lines_out = []
    for ln in text.split("\n"):
        if _hit(ln, exn):
            lines_out.append(f"<mark class='pv-hl' style='{_HL_STYLE}'>{_esc(ln)}</mark>")
        else:
            lines_out.append(_esc(ln))
    return ("<pre class='whitespace-pre-wrap text-xs overflow-auto max-h-[70vh] bg-gray-50 p-3 rounded'>"
            + "\n".join(lines_out) + "</pre>")
